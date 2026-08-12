"""CAPAG (Tesouro Nacional): nota oficial de capacidade de pagamento do município.

A CAPAG é publicada como planilha no CKAN do Tesouro Transparente. O coletor
descobre o arquivo mais recente via API CKAN, baixa o XLSX uma vez e mantém um
cache local em JSON indexado pelo código IBGE (renovado a cada 7 dias).
"""
import io
import json
import re
import time

import httpx
from openpyxl import load_workbook

from ..config import CAPAG_CACHE
from .base import get_json, resultado

CKAN_PACKAGE = "https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show"
CACHE_TTL = 7 * 24 * 3600  # 7 dias


def _achar_coluna(cabecalho: list[str], *padroes: str) -> int | None:
    for i, nome in enumerate(cabecalho):
        texto = str(nome or "").lower()
        if all(re.search(p, texto) for p in padroes):
            return i
    return None


def _parse_xlsx(conteudo: bytes) -> dict[str, dict]:
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb.active  # primeira aba: "Prévia da CAPAG" (classificação mais recente)
    linhas = ws.iter_rows(values_only=True)

    # localiza a linha de cabeçalho ("Código Município Completo" | Nome | UF | CAPAG | Notas 1-3)
    cabecalho = None
    for linha in linhas:
        valores = [str(v or "") for v in linha]
        if any("município completo" in v.lower() or "municipio completo" in v.lower() for v in valores):
            cabecalho = valores
            break
    if cabecalho is None:
        raise ValueError("cabeçalho 'Código Município Completo' não encontrado na planilha CAPAG")

    col_ibge = _achar_coluna(cabecalho, r"munic[ií]pio completo")
    col_capag = _achar_coluna(cabecalho, r"^capag$")
    col_endivid = _achar_coluna(cabecalho, r"^nota 1$")   # indicador 1: endividamento
    col_poupanca = _achar_coluna(cabecalho, r"^nota 2$")  # indicador 2: poupança corrente
    col_liquidez = _achar_coluna(cabecalho, r"^nota 3$")  # indicador 3: liquidez

    dados: dict[str, dict] = {}
    for linha in linhas:
        if linha is None or col_ibge is None or col_ibge >= len(linha) or linha[col_ibge] is None:
            continue
        cod = re.sub(r"\D", "", str(linha[col_ibge]))
        if len(cod) < 6:
            continue
        pega = lambda idx: (str(linha[idx]).strip() if idx is not None and linha[idx] is not None else None)
        dados[cod] = {
            "nota_capag": pega(col_capag),
            "nota_endividamento": pega(col_endivid),
            "nota_poupanca_corrente": pega(col_poupanca),
            "nota_liquidez": pega(col_liquidez),
        }
    wb.close()
    return dados


async def _carregar_base(client: httpx.AsyncClient) -> dict:
    if CAPAG_CACHE.exists():
        cache = json.loads(CAPAG_CACHE.read_text(encoding="utf-8"))
        if time.time() - cache.get("baixado_em", 0) < CACHE_TTL:
            return cache

    pacote = await get_json(client, CKAN_PACKAGE, params={"id": "capag-municipios"})
    recursos = [
        r
        for r in pacote["result"]["resources"]
        if r.get("format", "").upper() == "XLSX" and "metadado" not in r["name"].lower()
    ]
    if not recursos:
        raise ValueError("nenhum recurso XLSX encontrado no dataset CAPAG")
    recurso = recursos[-1]  # o dataset é ordenado cronologicamente; o último é o mais recente

    arquivo = await client.get(recurso["url"])
    arquivo.raise_for_status()
    municipios = _parse_xlsx(arquivo.content)

    cache = {
        "baixado_em": time.time(),
        "referencia": recurso["name"].strip(),
        "url": recurso["url"],
        "municipios": municipios,
    }
    CAPAG_CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
    return cache


async def coletar_capag(client: httpx.AsyncClient, cod_ibge: str) -> dict:
    try:
        base = await _carregar_base(client)
        registro = base["municipios"].get(str(cod_ibge))
        if registro is None:
            return resultado("capag", erro=f"município {cod_ibge} não consta na base CAPAG ({base['referencia']})")
        return resultado("capag", {**registro, "referencia": base["referencia"]})
    except Exception as exc:  # noqa: BLE001
        return resultado("capag", erro=f"{type(exc).__name__}: {exc}")
